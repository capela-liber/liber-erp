# -*- coding: utf-8 -*-
"""Reading Metabooks' answer to a shipment.

The delivery has no synchronous reply: a consolidated e-mail arrives after
midnight with an .xlsx report, and until it is read a batch says "sent" whether
or not Metabooks kept anything. The first real one, on 2026-09-04, said 229 of
278 titles went in and 49 came back -- and nothing in Odoo knew it.

Their refusal is of the whole record. A book refused over one column did not
lose that column: it lost the keywords, the Thema and the BISAC that travelled
with it. So reading the report is not bookkeeping, it is what puts those books
back in the queue.
"""

import base64
import io
import logging

from odoo import _, api, fields, models
from odoo.exceptions import UserError

_logger = logging.getLogger(__name__)

# The sheets that mean "this title did not go in". Their report also carries
# price notices and media rejections, which do not hold a title back.
SHEET_REJECTED = 'Rejeitados'
SHEET_PENDING = 'Em análise'

# Header names as their report writes them, mapped to what we do with them.
COL_GTIN = 'GTIN'
COL_MESSAGE = 'Mensagem de erro'
COL_CODE = 'Código do erro'
COL_FILE = 'Nome do arquivo'


class MetabooksImportReport(models.TransientModel):
    _name = 'metabooks.import.report'
    _description = 'Read a Metabooks upload report'

    file = fields.Binary('Report (.xlsx)', required=True)
    filename = fields.Char()
    batch_id = fields.Many2one('metabooks.export.batch', string='Batch')

    def action_import(self):
        self.ensure_one()
        rows = self._read(base64.b64decode(self.file))
        if not rows:
            raise UserError(_(
                "This report lists no refused title. Nothing to do -- which "
                "is the good outcome: every book in it was accepted."))

        recusados, orfaos = self._apply(rows)
        lotes = recusados.batch_id
        for lote in lotes:
            suas = recusados.filtered(lambda l: l.batch_id == lote)
            lote.report_on = fields.Datetime.now()
            lote.message_post(body=_(
                "Metabooks report read: %(n)s book(s) refused and put back in "
                "the queue. %(codes)s",
                n=len(suas),
                codes=', '.join(sorted({
                    '%s (%s)' % (l.error_code or '?', l.error_message or '')
                    for l in suas}))))

        if orfaos:
            _logger.warning(
                "Metabooks report: %s GTIN(s) not found in any sent batch: %s",
                len(orfaos), ', '.join(sorted(orfaos)[:20]))
        return self._done(recusados, orfaos)

    # ------------------------------------------------------------------ #

    @staticmethod
    def _read(data):
        """The refused rows of the report, as dicts keyed by their headers."""
        try:
            import openpyxl
            book = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
        except Exception as exc:  # noqa: BLE001 -- any bad file, one message
            raise UserError(_(
                "This does not look like a Metabooks report (.xlsx): %s", exc))

        rows = []
        for name in (SHEET_REJECTED, SHEET_PENDING):
            if name not in book.sheetnames:
                continue
            sheet = book[name]
            linhas = sheet.iter_rows(values_only=True)
            cabecalho = next(linhas, None)
            if not cabecalho:
                continue
            # By header, not by position: their layout differs between the
            # sheets, and a column added upstream must not shift our reading.
            indice = {str(c).strip(): i for i, c in enumerate(cabecalho) if c}
            if COL_GTIN not in indice:
                continue
            for linha in linhas:
                gtin = linha[indice[COL_GTIN]]
                if not gtin:
                    continue
                rows.append({
                    'gtin': str(gtin).strip().split('.')[0],
                    'message': MetabooksImportReport._cell(linha, indice, COL_MESSAGE)
                               or _("Held for internal review at Metabooks."),
                    'code': MetabooksImportReport._cell(linha, indice, COL_CODE),
                    'file': MetabooksImportReport._cell(linha, indice, COL_FILE),
                })
        return rows

    @staticmethod
    def _cell(linha, indice, header):
        if header not in indice:
            return ''
        valor = linha[indice[header]]
        return '' if valor is None else str(valor).strip()

    def _apply(self, rows):
        Line = self.env['metabooks.export.line']
        recusados = Line.browse()
        orfaos = set()
        for row in rows:
            linha = self._match(row)
            if not linha:
                orfaos.add(row['gtin'])
                continue
            linha._refuse(row['code'], row['message'])
            recusados |= linha
        return recusados, orfaos

    def _match(self, row):
        """The line this refusal belongs to.

        Their file name is ours with a user prefix and a hash suffix
        (!jsallum!_V_BR0089701_20260903_Alteracoes_3789aa...), so the batch is
        found by ours being inside theirs. Failing that -- an older report, a
        renamed file -- the most recent sent line for that ISBN is the honest
        second guess.
        """
        Line = self.env['metabooks.export.line']
        dominio = [('gtin', '=', row['gtin']),
                   ('batch_id.state', '=', 'sent')]
        if self.batch_id:
            dominio.append(('batch_id', '=', self.batch_id.id))

        candidatas = Line.search(dominio, order='id desc')
        nome_deles = row.get('file') or ''
        for linha in candidatas:
            nosso = (linha.batch_id.filename or '').rsplit('.', 1)[0]
            if nosso and nosso in nome_deles:
                return linha
        return candidatas[:1]

    def _done(self, recusados, orfaos):
        corpo = _("%s book(s) refused by Metabooks are back in the queue.",
                  len(recusados))
        if orfaos:
            corpo += _(" %s ISBN(s) in the report matched no sent batch.",
                       len(orfaos))
        return {
            'type': 'ir.actions.client',
            'tag': 'display_notification',
            'params': {
                'title': _("Metabooks report read"),
                'message': corpo,
                'type': 'warning' if orfaos else 'success',
                'sticky': bool(orfaos),
                'next': {'type': 'ir.actions.act_window_close'},
            },
        }
